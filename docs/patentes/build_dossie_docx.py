from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


BASE_DIR = Path(__file__).resolve().parent
SOURCE_MD = BASE_DIR / "dossie-tecnico-algoritmos-tachegando.md"
OUTPUT_DOCX = BASE_DIR / "dossie-tecnico-algoritmos-tachegando.docx"

BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
INK = RGBColor(11, 37, 69)
MUTED = RGBColor(90, 96, 105)
LIGHT_FILL = "F4F6F9"
HEADER_FILL = "F2F4F7"
BORDER = "D0D7DE"
WHITE = RGBColor(255, 255, 255)


ACCENT_REPLACEMENTS = {
    "Dossie": "Dossiê",
    "dossie": "dossiê",
    "tecnico": "técnico",
    "tecnica": "técnica",
    "tecnicos": "técnicos",
    "tecnicas": "técnicas",
    "metodo": "método",
    "metodos": "métodos",
    "nao": "não",
    "juridico": "jurídico",
    "juridica": "jurídica",
    "juridicas": "jurídicas",
    "estrategia": "estratégia",
    "avaliacao": "avaliação",
    "protecao": "proteção",
    "nucleo": "núcleo",
    "area": "área",
    "Area": "Área",
    "geografica": "geográfica",
    "geografico": "geográfico",
    "proximas": "próximas",
    "proximos": "próximos",
    "proxima": "próxima",
    "proximo": "próximo",
    "intencao": "intenção",
    "intencoes": "intenções",
    "Intencao": "Intenção",
    "Intencoes": "Intenções",
    "Requisicao": "Requisição",
    "requisicao": "requisição",
    "localizacao": "localização",
    "pontuacao": "pontuação",
    "distancia": "distância",
    "codigo": "código",
    "poligono": "polígono",
    "dinamica": "dinâmica",
    "variacao": "variação",
    "variacoes": "variações",
    "conclusao": "conclusão",
    "implementacao": "implementação",
    "atribuicao": "atribuição",
    "orquestracao": "orquestração",
    "expiracao": "expiração",
    "operacao": "operação",
    "analise": "análise",
    "evidencia": "evidência",
    "evidencias": "evidências",
    "funcoes": "funções",
    "restricao": "restrição",
    "decisao": "decisão",
    "custo-beneficio": "custo-benefício",
    "geodesica": "geodésica",
    "possiveis": "possíveis",
    "Possiveis": "Possíveis",
    "reivindicacao": "reivindicação",
    "reivindicacoes": "reivindicações",
    "saida": "saída",
    "Saida": "Saída",
    "relacao": "relação",
    "minimo": "mínimo",
    "modulo": "módulo",
    "modulos": "módulos",
    "Ja": "Já",
    "basica": "básica",
    "rotulo": "rótulo",
    "ate": "até",
    "ha": "há",
    "Ha": "Há",
    "generico": "genérico",
    "viavel": "viável",
    "viaveis": "viáveis",
    "maxima": "máxima",
    "tolerancia": "tolerância",
    "influencia": "influência",
    "praticos": "práticos",
    "oportunisticos": "oportunísticos",
    "solicitacao": "solicitação",
    "multiplas": "múltiplas",
    "configuravel": "configurável",
    "temporarias": "temporárias",
    "filtra-las": "filtrá-las",
    "intersecao": "interseção",
    "frequencia": "frequência",
    "historica": "histórica",
    "expansao": "expansão",
    "Persistencia": "Persistência",
    "Criacao": "Criação",
    "Geracao": "Geração",
    "Metodo": "Método",
    "Variacao": "Variação",
    "Mantem": "Mantém",
    "mantem": "mantém",
    "Ta Chegando": "Tá Chegando",
}


def set_run_font(run, name="Calibri", size=None, color=None, bold=None, italic=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def accent_plain(text: str) -> str:
    for source, target in ACCENT_REPLACEMENTS.items():
        text = re.sub(rf"\b{re.escape(source)}\b", target, text)
    text = text.replace(" e o **ERR", " é o **ERR")
    text = text.replace("projeto e o ", "projeto é o ")
    text = text.replace("atual e ", "atual é ")
    text = text.replace("nao e ", "não é ")
    text = text.replace(" e comum", " é comum")
    text = text.replace("O melhor caminho e ", "O melhor caminho é ")
    text = text.replace("A tese deve ", "A tese deve ")
    return text


def accent_outside_code(text: str) -> str:
    parts = re.split(r"(`[^`]+`)", text)
    return "".join(part if part.startswith("`") else accent_plain(part) for part in parts)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.start_type = WD_SECTION_START.NEW_PAGE
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    pf = normal.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(6)
    pf.line_spacing = 1.10

    for style_name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ]:
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.10

    for style_name in ["List Bullet", "List Number"]:
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(11)
        style.paragraph_format.space_after = Pt(8)
        style.paragraph_format.line_spacing = 1.167

    header_p = section.header.paragraphs[0]
    header_p.text = ""
    header_p.paragraph_format.space_after = Pt(0)
    run = header_p.add_run("Tá Chegando | Dossiê técnico de algoritmos")
    set_run_font(run, size=9, color=MUTED)
    add_bottom_border(header_p, BORDER, "4")

    footer_p = section.footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer_p.paragraph_format.space_before = Pt(0)
    footer_p.paragraph_format.space_after = Pt(0)
    run = footer_p.add_run("Confidencial - uso jurídico")
    set_run_font(run, size=9, color=MUTED)


def add_bottom_border(paragraph, color: str, size: str = "8") -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin_name, value in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn(f"w:{margin_name}"))
        if node is None:
            node = OxmlElement(f"w:{margin_name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_dxa: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_width(table, width_dxa: int = 9360, indent_dxa: int = 120) -> None:
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(width_dxa))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")


def repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_masthead(doc: Document) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run("DOCUMENTO TÉCNICO")
    set_run_font(run, size=10, color=BLUE, bold=True)

    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(4)
    run = title.add_run("Dossiê técnico de algoritmos - Tá Chegando")
    set_run_font(run, size=24, color=INK, bold=True)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(14)
    run = subtitle.add_run("ERR - Expansor de Rotas de Requisição e módulos relacionados")
    set_run_font(run, size=13, color=MUTED)

    rows = [
        ("Para:", "Advogado de propriedade intelectual"),
        ("De:", "Projeto Tá Chegando"),
        ("Data:", "30 de maio de 2026"),
        ("Finalidade:", "Base técnica para avaliação de proteção jurídica"),
        ("Status:", "Documento técnico; não substitui busca de anterioridade nem parecer jurídico"),
    ]
    for label, value in rows:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        label_run = p.add_run(label + " ")
        set_run_font(label_run, size=10.5, bold=True, color=INK)
        value_run = p.add_run(value)
        set_run_font(value_run, size=10.5, color=RGBColor(0, 0, 0))

    rule = doc.add_paragraph()
    rule.paragraph_format.space_before = Pt(10)
    rule.paragraph_format.space_after = Pt(12)
    add_bottom_border(rule, "9AA6B2", "8")


def add_markdown_runs(paragraph, text: str, size=None, color=None, italic=False) -> None:
    text = accent_outside_code(text)
    tokens = re.split(r"(`[^`]+`|\*\*[^*]+\*\*)", text)
    for token in tokens:
        if not token:
            continue
        if token.startswith("`") and token.endswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, name="Consolas", size=size or 9.2, color=DARK_BLUE)
        elif token.startswith("**") and token.endswith("**"):
            run = paragraph.add_run(accent_plain(token[2:-2]))
            set_run_font(run, size=size, color=color, bold=True, italic=italic)
        else:
            run = paragraph.add_run(token)
            set_run_font(run, size=size, color=color, italic=italic)


def add_body_paragraph(doc: Document, text: str, style=None) -> None:
    p = doc.add_paragraph(style=style)
    add_markdown_runs(p, text)


def add_quote_callout(doc: Document, text: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_width(table, 9360, 120)
    cell = table.cell(0, 0)
    shade_cell(cell, LIGHT_FILL)
    set_cell_margins(cell, 140, 180, 140, 180)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    add_markdown_runs(p, text, size=10.8, color=INK, italic=True)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def make_path_readable(text: str) -> str:
    if "/" in text and len(text) > 42:
        return text.replace("/", "/\n")
    return text


def add_markdown_table(doc: Document, table_lines: list[str]) -> None:
    rows = []
    for line in table_lines:
        parts = [part.strip() for part in line.strip().strip("|").split("|")]
        if all(re.fullmatch(r":?-{3,}:?", part or "") for part in parts):
            continue
        rows.append(parts)
    if not rows:
        return

    col_count = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=col_count)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    set_table_width(table, 9360, 120)

    widths = [4100, 4550, 710] if col_count == 3 else [int(9360 / col_count)] * col_count
    for row_idx, row in enumerate(rows):
        for col_idx in range(col_count):
            cell = table.cell(row_idx, col_idx)
            set_cell_width(cell, widths[min(col_idx, len(widths) - 1)])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_idx == 0:
                shade_cell(cell, HEADER_FILL)
            text = row[col_idx] if col_idx < len(row) else ""
            text = make_path_readable(text)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            if row_idx == 0:
                add_markdown_runs(p, text, size=9.5, color=INK)
                for run in p.runs:
                    run.bold = True
            else:
                add_markdown_runs(p, text, size=8.3 if col_idx == 1 else 8.8)
    repeat_table_header(table.rows[0])
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(4)


def add_mermaid_as_steps(doc: Document, block: str) -> None:
    labels = re.findall(r'\["([^"]+)"\]', block)
    seen: set[str] = set()
    ordered = []
    for label in labels:
        if label not in seen:
            seen.add(label)
            ordered.append(label)
    if not ordered:
        add_body_paragraph(doc, block)
        return
    for label in ordered:
        p = doc.add_paragraph(style="List Number")
        add_markdown_runs(p, label)


def parse_markdown(doc: Document, markdown: str) -> None:
    lines = markdown.splitlines()
    i = 0
    skipping_title = True
    while i < len(lines):
        raw = lines[i]
        line = raw.rstrip()
        if not line.strip():
            i += 1
            continue

        if skipping_title and line.startswith("# "):
            skipping_title = False
            i += 1
            continue
        if line.startswith("Data:"):
            i += 1
            continue

        if line.startswith("```mermaid"):
            block = []
            i += 1
            while i < len(lines) and not lines[i].startswith("```"):
                block.append(lines[i])
                i += 1
            add_mermaid_as_steps(doc, "\n".join(block))
            i += 1
            continue
        if line.startswith("```"):
            block = []
            i += 1
            while i < len(lines) and not lines[i].startswith("```"):
                block.append(lines[i])
                i += 1
            p = doc.add_paragraph()
            for idx, block_line in enumerate(block):
                if idx:
                    p.add_run().add_break()
                run = p.add_run(block_line)
                set_run_font(run, name="Consolas", size=9.2, color=DARK_BLUE)
            i += 1
            continue

        if line.startswith("|"):
            table_lines = []
            while i < len(lines) and lines[i].startswith("|"):
                table_lines.append(lines[i])
                i += 1
            add_markdown_table(doc, table_lines)
            continue

        if line.startswith("## "):
            p = doc.add_paragraph(style="Heading 1")
            add_markdown_runs(p, line[3:])
            i += 1
            continue
        if line.startswith("### "):
            p = doc.add_paragraph(style="Heading 2")
            add_markdown_runs(p, line[4:])
            i += 1
            continue
        if line.startswith("#### "):
            p = doc.add_paragraph(style="Heading 3")
            add_markdown_runs(p, line[5:])
            i += 1
            continue
        if line.startswith("> "):
            add_quote_callout(doc, line[2:])
            i += 1
            continue
        if line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            add_markdown_runs(p, line[2:])
            i += 1
            continue
        numbered = re.match(r"^(\d+)\.\s+(.*)$", line)
        if numbered:
            p = doc.add_paragraph(style="List Number")
            add_markdown_runs(p, numbered.group(2))
            i += 1
            continue

        add_body_paragraph(doc, line)
        i += 1


def main() -> int:
    if not SOURCE_MD.exists():
        print(f"Source not found: {SOURCE_MD}", file=sys.stderr)
        return 1

    doc = Document()
    configure_document(doc)
    add_masthead(doc)
    parse_markdown(doc, SOURCE_MD.read_text(encoding="utf-8"))
    doc.core_properties.title = "Dossiê técnico de algoritmos - Tá Chegando"
    doc.core_properties.subject = "ERR - Expansor de Rotas de Requisição"
    doc.core_properties.author = "Projeto Tá Chegando"
    doc.core_properties.comments = "Documento técnico para avaliação jurídica de algoritmos."
    doc.save(OUTPUT_DOCX)
    print(OUTPUT_DOCX)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
